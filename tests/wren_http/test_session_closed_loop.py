import asyncio
import json
from io import BytesIO

import duckdb
import httpx
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import SecretStr
from wren_http.bootstrap import create_app as create_wren_app
from wren_http.settings import ServiceSettings

from app.bootstrap import create_app
from app.infrastructure.wren_http import (
    WrenAdminSettings,
    WrenPublicationClient,
    WrenQuestionGateway,
    WrenQuestionGatewaySettings,
)
from app.modules.datasets import DatasetPublicationService, WorkbookUpload
from app.settings import Settings
from tests.test_sessions import ALICE, ask, create_session
from tests.wren_http.test_question_closed_loop import (
    count_agent_model,
    workbook_bytes,
)


def sources(candidate):
    return tuple(
        {
            "model": table["model"],
            "dataset_id": dataset.dataset_id,
            "source_name": dataset.source_name,
            "content_hash": dataset.content_hash,
            "sheet": table["sheet"],
        }
        for dataset in candidate.datasets
        for table in dataset.tables
    )


def register(registrar, candidate, token):
    asyncio.run(
        registrar.register_duckdb(
            workspace_id=candidate.workspace_id,
            publication_id=candidate.publication_id,
            workspace_token=token,
            project_relative_path=candidate.project_relative_path.as_posix(),
            sources=sources(candidate),
        )
    )


def test_agentscope_session_calls_wren_agent_and_tracks_new_publication(tmp_path):
    workspace_id = "session-closed-loop"
    admin_token = SecretStr("session-admin-token")
    workspace_token = SecretStr("session-workspace-token")
    publisher = DatasetPublicationService(tmp_path / "datasets")
    first = publisher.prepare(
        workspace_id,
        (
            WorkbookUpload(
                filename="guests.xlsx",
                dataset_id="test",
                content=workbook_bytes(),
            ),
        ),
    )
    model_calls = []
    wren_app = create_wren_app(
        ServiceSettings(
            admin_token=admin_token,
            workspaces_file=tmp_path / "wren" / "workspaces.json",
            duckdb_root=publisher.published_root,
            memory_enabled=False,
        ),
        model_factory=lambda _: count_agent_model(model_calls),
    )
    transport = httpx.ASGITransport(app=wren_app)
    registrar = WrenPublicationClient(
        WrenAdminSettings(base_url="http://wren.test", admin_token=admin_token),
        transport=transport,
    )
    register(registrar, first, workspace_token)
    publisher.activate(first)
    gateway = WrenQuestionGateway(
        WrenQuestionGatewaySettings(
            base_url="http://wren.test",
            service_token=workspace_token,
        ),
        transport=transport,
    )
    users = tmp_path / "users.json"
    users.write_text(
        json.dumps(
            [
                {
                    "user_id": "alice",
                    "api_token": "alice-session-test",
                    "workspace_ids": [workspace_id],
                }
            ]
        ),
        encoding="utf-8",
    )
    settings = Settings(
        _env_file=None,
        api_users_path=users,
        workspace_id=workspace_id,
        control_database=tmp_path / "control.duckdb",
        dataset_root=tmp_path / "datasets",
        session_workspace=tmp_path / "session-workspaces",
    )

    with TestClient(
        create_app(
            settings,
            dataset_publisher=publisher,
            question_gateway=gateway,
        ),
        headers=ALICE,
    ) as client:
        agent_id, session_id = create_session(client)
        first_reply, _ = ask(client, agent_id, session_id, "一共有多少条记录？", "turn-1")
        second_reply, _ = ask(client, agent_id, session_id, "再确认一次总数", "turn-2")

    assert first_reply["result"]["rows"] == [{"total": 2}]
    assert second_reply["result"]["rows"] == [{"total": 2}]
    assert first_reply["result"]["publication_id"] == first.publication_id
    assert [item["name"] for item in second_reply["result"]["tool_trace"]] == [
        "wren_list_models",
        "wren_query",
    ]
    assert first_reply["result"]["sources"][0]["source_name"] == "guests.xlsx"

    updated_workbook = workbook_bytes()
    book = load_workbook(BytesIO(updated_workbook))
    book.active.append(["Q003", "Three"])
    output = BytesIO()
    book.save(output)
    book.close()
    second = publisher.prepare(
        workspace_id,
        (
            WorkbookUpload(
                filename="guests.xlsx",
                dataset_id="test",
                content=output.getvalue(),
            ),
        ),
    )
    register(registrar, second, workspace_token)
    publisher.activate(second)

    with TestClient(
        create_app(
            settings,
            dataset_publisher=publisher,
            question_gateway=gateway,
        ),
        headers=ALICE,
    ) as client:
        third_reply, _ = ask(client, agent_id, session_id, "更新后有多少条？", "turn-3")

    assert third_reply["result"]["rows"] == [{"total": 3}]
    assert third_reply["result"]["publication_id"] == second.publication_id
    with duckdb.connect(str(second.database_path), read_only=True) as connection:
        expected = connection.execute("SELECT COUNT(*) FROM t_test_sheet_1").fetchone()[0]
    assert expected == third_reply["result"]["rows"][0]["total"]
    assert len(model_calls) == 9
