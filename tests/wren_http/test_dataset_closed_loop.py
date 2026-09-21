import asyncio
import json
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import duckdb
import httpx
import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pydantic import SecretStr
from wren_http.bootstrap import create_app as create_wren_app
from wren_http.settings import ServiceSettings

from app.contracts.base import AppError
from app.infrastructure.auth import IdentityProvider, UserCredential
from app.infrastructure.wren_http import (
    WrenAdminSettings,
    WrenHttpClient,
    WrenHttpSettings,
    WrenPublicationClient,
)
from app.modules.datasets import DatasetPublicationService, RelationshipInput, WorkbookUpload
from app.modules.datasets.api import build_dataset_router


def workbook_bytes(headers, rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def workbook_bytes_with_stale_dimension(headers, rows):
    source = BytesIO(workbook_bytes(headers, rows))
    output = BytesIO()
    with ZipFile(source) as incoming, ZipFile(output, "w", ZIP_DEFLATED) as outgoing:
        for item in incoming.infolist():
            content = incoming.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b'ref="A1:B3"', b'ref="A1:A1"', 1)
            outgoing.writestr(item, content)
    return output.getvalue()


def model_for(candidate, filename):
    dataset = next(item for item in candidate.datasets if item.source_name == filename)
    return dataset.tables[0]["model"]


async def register(registrar, candidate, token):
    await registrar.register_duckdb(
        workspace_id=candidate.workspace_id,
        publication_id=candidate.publication_id,
        workspace_token=token,
        project_relative_path=candidate.project_relative_path.as_posix(),
    )


def test_upload_update_register_query_and_failure_keep_previous_publication(tmp_path):
    workspace = "conference-a"
    admin_token = SecretStr("closed-loop-admin-token")
    query_token = SecretStr("closed-loop-query-token")
    publisher = DatasetPublicationService(tmp_path / "datasets")
    wren_app = create_wren_app(
        ServiceSettings(
            admin_token=admin_token,
            workspaces_file=tmp_path / "wren" / "workspaces.json",
            duckdb_root=publisher.published_root,
        )
    )
    transport = httpx.ASGITransport(app=wren_app)
    registrar = WrenPublicationClient(
        WrenAdminSettings(base_url="http://wren.test", admin_token=admin_token),
        transport=transport,
    )

    guests_v1 = workbook_bytes(
        ["guest_id", "name", "country"],
        [["A001", "Alice", "CN"], ["A002", "Bob", "US"]],
    )
    events = workbook_bytes(
        ["event_id", "title", "guest_id"],
        [["E01", "Opening", "A001"]],
    )
    first = publisher.prepare(
        workspace,
        (
            WorkbookUpload(filename="guests.xlsx", content=guests_v1),
            WorkbookUpload(filename="events.xlsx", content=events),
        ),
    )
    asyncio.run(register(registrar, first, query_token))
    first_result = publisher.activate(first)
    assert first_result.changed and len(first_result.datasets) == 2

    async def query(candidate, sql):
        client = WrenHttpClient(
            WrenHttpSettings(
                base_url="http://wren.test",
                service_token=query_token,
                workspace_id=workspace,
                publication_id=candidate.publication_id,
            ),
            transport=transport,
        )
        return await client.query(sql)

    guest_model_v1 = model_for(first, "guests.xlsx")
    event_model_v1 = model_for(first, "events.xlsx")
    first_count = asyncio.run(
        query(first, f'SELECT COUNT(*) AS n FROM "{guest_model_v1}"')
    )
    assert first_count.rows == ({"n": 2},)
    assert asyncio.run(
        query(first, f'SELECT COUNT(*) AS n FROM "{event_model_v1}"')
    ).rows == ({"n": 1},)

    guests_v2 = workbook_bytes(
        ["guest_id", "name", "country"],
        [
            ["A001", "Alice", "CN"],
            ["A002", "Bob", "US"],
            ["A003", "Carol", "FR"],
        ],
    )
    second = publisher.prepare(
        workspace, (WorkbookUpload(filename="guests.xlsx", content=guests_v2),)
    )
    assert second.publication_id != first.publication_id
    assert {item.dataset_id for item in second.datasets} == {
        item.dataset_id for item in first.datasets
    }
    asyncio.run(register(registrar, second, query_token))
    publisher.activate(second)
    guest_model_v2 = model_for(second, "guests.xlsx")
    event_model_v2 = model_for(second, "events.xlsx")
    assert asyncio.run(
        query(second, f'SELECT COUNT(*) AS n FROM "{guest_model_v2}"')
    ).rows == ({"n": 3},)
    assert asyncio.run(
        query(second, f'SELECT COUNT(*) AS n FROM "{event_model_v2}"')
    ).rows == ({"n": 1},)
    assert asyncio.run(
        query(first, f'SELECT COUNT(*) AS n FROM "{guest_model_v1}"')
    ).rows == ({"n": 2},)

    repeated = publisher.prepare(
        workspace, (WorkbookUpload(filename="guests.xlsx", content=guests_v2),)
    )
    assert not repeated.changed and repeated.publication_id == second.publication_id
    assert not publisher.activate(repeated).changed

    with pytest.raises(AppError) as caught:
        publisher.prepare(
            workspace,
            (WorkbookUpload(filename="guests.xlsx", content=b"not-an-xlsx"),),
        )
    assert caught.value.code == "EXCEL_INVALID"
    active = publisher.active(workspace)
    assert active is not None and active.publication_id == second.publication_id
    assert asyncio.run(
        query(second, f'SELECT COUNT(*) AS n FROM "{guest_model_v2}"')
    ).rows == ({"n": 3},)

    related = publisher.prepare_relationship(
        workspace,
        RelationshipInput(
            left_model=event_model_v2,
            left_column="guest_id",
            right_model=guest_model_v2,
            right_column="guest_id",
            join_type="MANY_TO_ONE",
        ),
    )
    asyncio.run(register(registrar, related, query_token))
    related_result = publisher.activate(related)
    assert related_result.relationships[0].join_type == "MANY_TO_ONE"
    joined = asyncio.run(
        query(
            related,
            f'SELECT COUNT(*) AS n FROM "{event_model_v2}" e '
            f'JOIN "{guest_model_v2}" g ON e.guest_id=g.guest_id',
        )
    )
    assert joined.rows == ({"n": 1},)
    target = json.loads((related.project_path / "target" / "mdl.json").read_text())
    assert target["relationships"][0]["joinType"] == "MANY_TO_ONE"

    event_dataset = next(
        item for item in related.datasets if item.source_name == "events.xlsx"
    )
    event_source = event_dataset.source_path
    removed = publisher.prepare_removal(workspace, event_dataset.dataset_id)
    asyncio.run(register(registrar, removed, query_token))
    removed_result = publisher.activate(removed)
    assert removed_result.changed
    assert [item.source_name for item in removed_result.datasets] == ["guests.xlsx"]
    assert event_source.is_file()
    assert first.database_path.is_file() and second.database_path.is_file()
    with pytest.raises(AppError):
        asyncio.run(query(removed, f'SELECT COUNT(*) AS n FROM "{event_model_v2}"'))
    assert asyncio.run(
        query(second, f'SELECT COUNT(*) AS n FROM "{event_model_v2}"')
    ).rows == ({"n": 1},)

    restored = publisher.rollback(workspace, second.publication_id)
    assert restored.changed
    assert {item.source_name for item in restored.datasets} == {
        "guests.xlsx",
        "events.xlsx",
    }
    assert asyncio.run(
        query(second, f'SELECT COUNT(*) AS n FROM "{event_model_v2}"')
    ).rows == ({"n": 1},)

    rolled_back = publisher.rollback(workspace, first.publication_id)
    assert rolled_back.changed
    assert rolled_back.publication_id == first.publication_id
    assert publisher.active(workspace).publication_id == first.publication_id
    assert asyncio.run(
        query(first, f'SELECT COUNT(*) AS n FROM "{guest_model_v1}"')
    ).rows == ({"n": 2},)


def test_same_named_excel_is_isolated_between_workspaces(tmp_path):
    admin_token = SecretStr("isolation-admin-token")
    token_a = SecretStr("conference-a-query-token")
    token_b = SecretStr("conference-b-query-token")
    publisher = DatasetPublicationService(tmp_path / "datasets")
    wren_app = create_wren_app(
        ServiceSettings(
            admin_token=admin_token,
            workspaces_file=tmp_path / "wren" / "workspaces.json",
            duckdb_root=publisher.published_root,
        )
    )
    transport = httpx.ASGITransport(app=wren_app)
    registrar = WrenPublicationClient(
        WrenAdminSettings(base_url="http://wren.test", admin_token=admin_token),
        transport=transport,
    )

    publications = {}
    for workspace, token, guest in (
        ("conference-isolation-a", token_a, "Only A"),
        ("conference-isolation-b", token_b, "Only B"),
    ):
        candidate = publisher.prepare(
            workspace,
            (
                WorkbookUpload(
                    filename="guests.xlsx",
                    content=workbook_bytes(["guest_id", "name"], [["G001", guest]]),
                ),
            ),
        )
        asyncio.run(register(registrar, candidate, token))
        publisher.activate(candidate)
        publications[workspace] = (candidate, token)

    async def names(workspace):
        candidate, token = publications[workspace]
        model = model_for(candidate, "guests.xlsx")
        client = WrenHttpClient(
            WrenHttpSettings(
                base_url="http://wren.test",
                service_token=token,
                workspace_id=workspace,
                publication_id=candidate.publication_id,
            ),
            transport=transport,
        )
        return await client.query(f'SELECT name FROM "{model}"')

    assert asyncio.run(names("conference-isolation-a")).rows == ({"name": "Only A"},)
    assert asyncio.run(names("conference-isolation-b")).rows == ({"name": "Only B"},)

    publication_a, _ = publications["conference-isolation-a"]
    removed_a = publisher.prepare_removal(
        "conference-isolation-a", publication_a.datasets[0].dataset_id
    )
    asyncio.run(register(registrar, removed_a, token_a))
    publisher.activate(removed_a)

    assert publisher.active("conference-isolation-a").datasets == ()
    assert publisher.active("conference-isolation-b").publication_id == publications[
        "conference-isolation-b"
    ][0].publication_id
    assert asyncio.run(names("conference-isolation-b")).rows == ({"name": "Only B"},)


def test_management_http_uploads_multiple_files_and_publishes(tmp_path):
    from types import SimpleNamespace

    workspace = "conference-http"
    admin_token = SecretStr("management-admin-token")
    query_token = SecretStr("management-query-token")
    publisher = DatasetPublicationService(tmp_path / "datasets")
    wren_app = create_wren_app(
        ServiceSettings(
            admin_token=admin_token,
            workspaces_file=tmp_path / "wren" / "workspaces.json",
            duckdb_root=publisher.published_root,
        )
    )
    transport = httpx.ASGITransport(app=wren_app)
    registrar = WrenPublicationClient(
        WrenAdminSettings(base_url="http://wren.test", admin_token=admin_token),
        transport=transport,
    )
    settings = SimpleNamespace(
        workspace_id=workspace,
        dataset_max_upload_bytes=5 * 1024 * 1024,
        wren_workspace_token=query_token,
    )

    async def current_user():
        return "tester"

    management = FastAPI()
    management.include_router(
        build_dataset_router(
            settings,
            current_user,
            publisher,
            registrar,
            IdentityProvider(
                workspace,
                (
                    UserCredential(
                        user_id="tester",
                        api_token=SecretStr("dataset-test-user-token"),
                        workspace_ids=(workspace,),
                    ),
                ),
            ),
        ),
        prefix="/api/v1",
    )
    guests = workbook_bytes(["guest_id", "name"], [["H001", "HTTP Guest"]])
    events = workbook_bytes(
        ["event_id", "title", "guest_id"],
        [["HE01", "HTTP Event", "H001"]],
    )
    files = [
        (
            "files",
            (
                "guests.xlsx",
                guests,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        ),
        (
            "files",
            (
                "events.xlsx",
                events,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        ),
    ]
    with TestClient(management) as client:
        response = client.post(
            f"/api/v1/workspaces/{workspace}/uploads", files=files
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["changed"] and len(body["datasets"]) == 2
        guest_model = next(
            item["tables"][0]["model"]
            for item in body["datasets"]
            if item["source_name"] == "guests.xlsx"
        )
        event_model = next(
            item["tables"][0]["model"]
            for item in body["datasets"]
            if item["source_name"] == "events.xlsx"
        )
        relationship = client.post(
            f"/api/v1/workspaces/{workspace}/relationships",
            json={
                "left_model": event_model,
                "left_column": "guest_id",
                "right_model": guest_model,
                "right_column": "guest_id",
                "join_type": "MANY_TO_ONE",
            },
        )
        assert relationship.status_code == 200, relationship.text
        relationship_body = relationship.json()
        assert relationship_body["relationships"][0]["join_type"] == "MANY_TO_ONE"
        active = client.get(
            f"/api/v1/workspaces/{workspace}/datasets"
        ).json()
        assert active["publication_id"] == relationship_body["publication_id"]
        repeated = client.post(
            f"/api/v1/workspaces/{workspace}/uploads", files=files
        ).json()
        assert not repeated["changed"]
        updated = client.post(
            f"/api/v1/workspaces/{workspace}/uploads",
            files={
                "files": (
                    "guests.xlsx",
                    workbook_bytes(
                        ["guest_id", "name"],
                        [["H001", "HTTP Guest"], ["H002", "New Guest"]],
                    ),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        ).json()
        assert updated["publication_id"] != body["publication_id"]
        events_id = next(
            item["dataset_id"]
            for item in updated["datasets"]
            if item["source_name"] == "events.xlsx"
        )
        removed = client.delete(
            f"/api/v1/workspaces/{workspace}/datasets/{events_id}"
        )
        assert removed.status_code == 200, removed.text
        removed_body = removed.json()
        assert [item["source_name"] for item in removed_body["datasets"]] == [
            "guests.xlsx"
        ]
        publications = client.get(
            f"/api/v1/workspaces/{workspace}/publications"
        ).json()
        assert publications[0]["active"]
        assert publications[0]["publication_id"] == removed_body["publication_id"]
        assert len(publications) == 4
        rollback = client.post(
            f"/api/v1/workspaces/{workspace}/publications/"
            f"{body['publication_id']}/activate"
        ).json()
        assert rollback["changed"]
        assert rollback["publication_id"] == body["publication_id"]
        active = client.get(
            f"/api/v1/workspaces/{workspace}/datasets"
        ).json()
        assert active["publication_id"] == body["publication_id"]

    guest_model = next(
        item["tables"][0]["model"]
        for item in body["datasets"]
        if item["source_name"] == "guests.xlsx"
    )
    query_client = WrenHttpClient(
        WrenHttpSettings(
            base_url="http://wren.test",
            service_token=query_token,
            workspace_id=workspace,
            publication_id=body["publication_id"],
        ),
        transport=transport,
    )
    result = asyncio.run(
        query_client.query(f'SELECT COUNT(*) AS n FROM "{guest_model}"')
    )
    assert result.rows == ({"n": 1},)


def test_management_keeps_active_publication_when_wren_registration_fails(tmp_path):
    from types import SimpleNamespace

    workspace = "conference-rollback"
    query_token = SecretStr("rollback-query-token")
    publisher = DatasetPublicationService(tmp_path / "datasets")
    initial = publisher.prepare(
        workspace,
        (
            WorkbookUpload(
                filename="guests.xlsx",
                content=workbook_bytes(["guest_id"], [["R001"]]),
            ),
        ),
    )
    publisher.activate(initial)

    class FailingRegistrar:
        async def register_duckdb(self, **kwargs):
            raise AppError("WREN_UNAVAILABLE", "问数引擎暂时不可用，数据未发布。", 503)

    settings = SimpleNamespace(
        workspace_id=workspace,
        dataset_max_upload_bytes=5 * 1024 * 1024,
        wren_workspace_token=query_token,
    )

    async def current_user():
        return "tester"

    management = FastAPI()
    management.include_router(
        build_dataset_router(
            settings,
            current_user,
            publisher,
            FailingRegistrar(),
            IdentityProvider(
                workspace,
                (
                    UserCredential(
                        user_id="tester",
                        api_token=SecretStr("rollback-test-user-token"),
                        workspace_ids=(workspace,),
                    ),
                ),
            ),
        ),
        prefix="/api/v1",
    )
    changed = workbook_bytes(["guest_id"], [["R001"], ["R002"]])
    with TestClient(management) as client:
        response = client.post(
            f"/api/v1/workspaces/{workspace}/uploads",
            files={
                "files": (
                    "guests.xlsx",
                    changed,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
    assert response.status_code == 503
    assert publisher.active(workspace).publication_id == initial.publication_id


def test_import_streams_cells_when_xlsx_dimension_metadata_is_stale(tmp_path):
    publisher = DatasetPublicationService(tmp_path / "datasets")
    candidate = publisher.prepare(
        "conference-stale-dimension",
        (
            WorkbookUpload(
                filename="guests.xlsx",
                content=workbook_bytes_with_stale_dimension(
                    ["guest_id", "name"],
                    [["S001", "One"], ["S002", "Two"]],
                ),
            ),
        ),
    )
    table = candidate.datasets[0].tables[0]
    assert table["row_count"] == 2
    assert [column["name"] for column in table["columns"]] == ["guest_id", "name"]


def test_control_database_schema_upgrades_to_standard_wren_project(tmp_path):
    root = tmp_path / "datasets"
    project = root / "published" / "legacy-space" / "p_legacy"
    project.mkdir(parents=True)
    database = project / "data.duckdb"
    with duckdb.connect(str(database)) as connection:
        connection.execute("CREATE TABLE t_legacy_sheet_1(id BIGINT)")
        connection.execute("INSERT INTO t_legacy_sheet_1 VALUES (1)")
    control = root / "control.duckdb"
    snapshot = [
        {
            "dataset_id": "legacy",
            "source_name": "legacy.xlsx",
            "content_hash": "a" * 64,
            "source_path": str(root / "sources" / "legacy.xlsx"),
            "tables": [
                {
                    "sheet": "Data",
                    "model": "ds_legacy_sheet_1",
                    "table": "t_legacy_sheet_1",
                    "row_count": 1,
                    "columns": [{"name": "id", "type": "BIGINT"}],
                }
            ],
        }
    ]
    with duckdb.connect(str(control)) as connection:
        connection.execute(
            """
            CREATE TABLE dataset_publications(
                workspace_id VARCHAR NOT NULL,
                publication_id VARCHAR NOT NULL,
                database_path VARCHAR NOT NULL,
                obsolete_payload VARCHAR NOT NULL,
                signature VARCHAR NOT NULL,
                datasets_json VARCHAR,
                created_at VARCHAR NOT NULL,
                PRIMARY KEY(workspace_id, publication_id)
            )
            """
        )
        connection.execute(
            "INSERT INTO dataset_publications VALUES (?,?,?,?,?,?,?)",
            (
                "legacy-space",
                "p_legacy",
                str(database),
                "obsolete",
                "legacy-signature",
                json.dumps(snapshot),
                "2026-09-20T00:00:00+00:00",
            ),
        )
    DatasetPublicationService(root)
    with duckdb.connect(str(control), read_only=True) as connection:
        columns = {
            row[1]
            for row in connection.execute(
                "PRAGMA table_info(dataset_publications)"
            ).fetchall()
        }
    assert "obsolete_payload" not in columns
    assert "project_path" in columns
    assert (project / "wren_project.yml").is_file()
    assert (project / "models" / "ds_legacy_sheet_1" / "metadata.yml").is_file()
