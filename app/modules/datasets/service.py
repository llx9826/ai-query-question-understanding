from __future__ import annotations

import datetime as dt
import decimal
import hashlib
import json
import os
import re
import shutil
import uuid
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from typing import Any

import duckdb
from openpyxl import load_workbook

from app.contracts.base import AppError
from app.infrastructure.duckdb_control import control_database_lock

from .schemas import (
    DatasetRecord,
    DatasetSummary,
    PublicationCandidate,
    PublicationResult,
    PublicationSummary,
    RelationshipDefinition,
    RelationshipInput,
    WorkbookUpload,
)
from .wren_project import write_wren_project

_ID = re.compile(r"^[a-zA-Z0-9][a-zA-Z0-9._-]{0,127}$")
_PUBLISHER_VERSION = "excel-wren-project-v3"


def _quoted(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def _stable_dataset_id(workspace_id: str, filename: str) -> str:
    digest = hashlib.sha256(f"{workspace_id}\0{filename.casefold()}".encode()).hexdigest()[:20]
    return f"ds_{digest}"


def _headers(values: tuple[Any, ...]) -> tuple[str, ...]:
    last = max((index for index, value in enumerate(values) if value not in (None, "")), default=-1)
    if last < 0:
        raise AppError("EXCEL_HEADER_MISSING", "工作表首行必须包含列名。", 422)
    result, used = [], set()
    for index, raw in enumerate(values[: last + 1], start=1):
        base = str(raw).strip() if raw not in (None, "") else f"Column_{index}"
        candidate, suffix = base, 2
        while candidate.casefold() in used:
            candidate = f"{base}__{suffix}"
            suffix += 1
        used.add(candidate.casefold())
        result.append(candidate)
    return tuple(result)


def _kind(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "BOOLEAN"
    if isinstance(value, int):
        return "BIGINT"
    if isinstance(value, (float, decimal.Decimal)):
        return "DOUBLE"
    if isinstance(value, dt.datetime):
        return "TIMESTAMP"
    if isinstance(value, dt.date):
        return "DATE"
    return "VARCHAR"


def _merge_kind(current: str | None, incoming: str | None) -> str | None:
    if incoming is None or current == incoming:
        return current or incoming
    if {current, incoming} <= {"BIGINT", "DOUBLE"}:
        return "DOUBLE"
    return "VARCHAR"


def _duckdb_value(value: Any) -> Any:
    if isinstance(value, decimal.Decimal):
        return float(value)
    return value


def _one_dict(cursor) -> dict[str, Any] | None:
    row = cursor.fetchone()
    if row is None:
        return None
    columns = [item[0] for item in cursor.description]
    return dict(zip(columns, row, strict=True))


def _all_dicts(cursor) -> list[dict[str, Any]]:
    columns = [item[0] for item in cursor.description]
    return [dict(zip(columns, row, strict=True)) for row in cursor.fetchall()]


def _summaries(datasets: tuple[DatasetRecord, ...]) -> tuple[DatasetSummary, ...]:
    return tuple(
        DatasetSummary(
            dataset_id=dataset.dataset_id,
            source_name=dataset.source_name,
            content_hash=dataset.content_hash,
            tables=dataset.tables,
        )
        for dataset in datasets
    )


class DatasetPublicationService:
    def __init__(
        self,
        root: Path,
        *,
        control_database: Path | None = None,
        max_upload_bytes: int = 50 * 1024 * 1024,
        max_rows_per_sheet: int = 100_000,
        max_columns_per_sheet: int = 200,
        max_cells_per_workbook: int = 1_000_000,
    ):
        self.root = root.resolve()
        self.source_root = self.root / "sources"
        self.published_root = self.root / "published"
        self.metadata_path = (control_database or self.root / "control.duckdb").resolve()
        self.metadata_path.parent.mkdir(parents=True, exist_ok=True)
        self.max_upload_bytes = max_upload_bytes
        self.max_rows_per_sheet = max_rows_per_sheet
        self.max_columns_per_sheet = max_columns_per_sheet
        self.max_cells_per_workbook = max_cells_per_workbook
        self._lock = control_database_lock(self.metadata_path)
        self.source_root.mkdir(parents=True, exist_ok=True)
        self.published_root.mkdir(parents=True, exist_ok=True)
        self._initialize()

    @contextmanager
    def _connect(self):
        with self._lock:
            connection = duckdb.connect(str(self.metadata_path))
            try:
                connection.begin()
                yield connection
                connection.commit()
            except Exception:
                connection.rollback()
                raise
            finally:
                connection.close()

    def _initialize(self):
        with self._connect() as connection:
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS dataset_workspaces(
                    workspace_id VARCHAR PRIMARY KEY,
                    active_publication_id VARCHAR
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS datasets(
                    workspace_id VARCHAR NOT NULL,
                    dataset_id VARCHAR NOT NULL,
                    source_name VARCHAR NOT NULL,
                    content_hash VARCHAR NOT NULL,
                    source_path VARCHAR NOT NULL,
                    schema_json VARCHAR NOT NULL,
                    updated_at VARCHAR NOT NULL,
                    PRIMARY KEY(workspace_id, dataset_id),
                    UNIQUE(workspace_id, source_name)
                )
                """
            )
            connection.execute(
                """
                CREATE TABLE IF NOT EXISTS dataset_publications(
                    workspace_id VARCHAR NOT NULL,
                    publication_id VARCHAR NOT NULL,
                    database_path VARCHAR NOT NULL,
                    project_path VARCHAR NOT NULL,
                    signature VARCHAR NOT NULL,
                    datasets_json VARCHAR,
                    created_at VARCHAR NOT NULL,
                    PRIMARY KEY(workspace_id, publication_id)
                )
                """
            )
            columns = {
                row[1]
                for row in connection.execute(
                    "PRAGMA table_info(dataset_publications)"
                ).fetchall()
            }
            if "project_path" not in columns:
                connection.execute(
                    """
                    CREATE TABLE dataset_publications_v2(
                        workspace_id VARCHAR NOT NULL,
                        publication_id VARCHAR NOT NULL,
                        database_path VARCHAR NOT NULL,
                        project_path VARCHAR NOT NULL,
                        signature VARCHAR NOT NULL,
                        datasets_json VARCHAR,
                        created_at VARCHAR NOT NULL,
                        PRIMARY KEY(workspace_id, publication_id)
                    )
                    """
                )
                connection.execute(
                    """
                    INSERT INTO dataset_publications_v2
                    SELECT workspace_id, publication_id, database_path,
                           regexp_replace(database_path, '[\\\\/][^\\\\/]+$', ''),
                           signature, datasets_json, created_at
                    FROM dataset_publications
                    """
                )
                connection.execute("DROP TABLE dataset_publications")
                connection.execute(
                    "ALTER TABLE dataset_publications_v2 RENAME TO dataset_publications"
                )
                columns = {
                    row[1]
                    for row in connection.execute(
                        "PRAGMA table_info(dataset_publications)"
                    ).fetchall()
                }
            if "datasets_json" not in columns:
                connection.execute(
                    "ALTER TABLE dataset_publications ADD COLUMN datasets_json TEXT"
                )
            if "relationships_json" not in columns:
                connection.execute(
                    "ALTER TABLE dataset_publications ADD COLUMN relationships_json TEXT"
                )
            self._restore_missing_projects(connection)

    def _restore_missing_projects(self, connection) -> None:
        rows = _all_dicts(
            connection.execute(
                "SELECT workspace_id,publication_id,database_path,project_path,"
                "datasets_json,relationships_json "
                "FROM dataset_publications"
            )
        )
        for row in rows:
            project_path = Path(row["project_path"])
            if (project_path / "wren_project.yml").is_file() or not row["datasets_json"]:
                continue
            datasets = tuple(
                DatasetRecord.model_validate(item)
                for item in json.loads(row["datasets_json"])
            )
            models = [
                {
                    "name": table["model"],
                    "description": f"{dataset.source_name} / {table['sheet']}",
                    "tableReference": {
                        "schema": "main",
                        "table": table["table"],
                    },
                    "columns": table["columns"],
                }
                for dataset in datasets
                for table in dataset.tables
            ]
            if Path(row["database_path"]).is_file() and models:
                write_wren_project(
                    project_path,
                    workspace_id=row["workspace_id"],
                    publication_id=row["publication_id"],
                    models=models,
                    relationships=json.loads(row["relationships_json"] or "[]"),
                )

    def prepare(
        self, workspace_id: str, uploads: tuple[WorkbookUpload, ...]
    ) -> PublicationCandidate:
        if not _ID.fullmatch(workspace_id):
            raise AppError("WORKSPACE_INVALID", "数据空间标识不合法。", 422)
        if not uploads:
            raise AppError("UPLOAD_EMPTY", "至少需要上传一个 Excel 文件。", 422)
        names = [Path(item.filename).name.casefold() for item in uploads]
        if len(names) != len(set(names)):
            raise AppError("UPLOAD_DUPLICATE_NAME", "同一批次不能包含重名文件。", 422)

        with self._lock:
            existing, base_publication = self._current(workspace_id)
            existing_by_id = {item.dataset_id: item for item in existing}
            existing_by_name = {item.source_name.casefold(): item for item in existing}
            replacements: dict[str, DatasetRecord] = {}
            for upload in uploads:
                record = self._store_upload(workspace_id, upload)
                same_name = existing_by_name.get(record.source_name.casefold())
                if same_name and same_name.dataset_id != record.dataset_id:
                    raise AppError(
                        "DATASET_ID_CONFLICT",
                        "该文件名已属于另一个数据集，请使用原数据集标识更新。",
                        409,
                    )
                previous = existing_by_id.get(record.dataset_id)
                replacements[record.dataset_id] = (
                    previous
                    if previous and previous.content_hash == record.content_hash
                    else record
                )
            combined = {record.dataset_id: record for record in existing}
            combined.update(replacements)
            datasets = tuple(sorted(combined.values(), key=lambda value: value.dataset_id))
            relationships = self._relationships(workspace_id, base_publication)
            model_names = {
                table["model"] for dataset in datasets for table in dataset.tables
            }
            relationships = tuple(
                item
                for item in relationships
                if set(item.models).issubset(model_names)
            )
            return self._prepare_candidate(
                workspace_id, base_publication, datasets, relationships
            )

    def prepare_removal(
        self, workspace_id: str, dataset_id: str
    ) -> PublicationCandidate:
        if not _ID.fullmatch(workspace_id) or not _ID.fullmatch(dataset_id):
            raise AppError("DATASET_INVALID", "数据空间或数据集标识不合法。", 422)
        with self._lock:
            existing, base_publication = self._current(workspace_id)
            if not any(item.dataset_id == dataset_id for item in existing):
                raise AppError("DATASET_NOT_FOUND", "数据集不存在或已被移除。", 404)
            datasets = tuple(item for item in existing if item.dataset_id != dataset_id)
            model_names = {
                table["model"] for dataset in datasets for table in dataset.tables
            }
            relationships = tuple(
                item
                for item in self._relationships(workspace_id, base_publication)
                if set(item.models).issubset(model_names)
            )
            return self._prepare_candidate(
                workspace_id, base_publication, datasets, relationships
            )

    def prepare_relationship(
        self, workspace_id: str, request: RelationshipInput
    ) -> PublicationCandidate:
        if not _ID.fullmatch(workspace_id):
            raise AppError("WORKSPACE_INVALID", "数据空间标识不合法。", 422)
        with self._lock:
            datasets, base_publication = self._current(workspace_id)
            if base_publication is None:
                raise AppError("PUBLICATION_NOT_FOUND", "当前空间尚未发布数据。", 409)
            tables = {
                table["model"]: table
                for dataset in datasets
                for table in dataset.tables
            }
            left = tables.get(request.left_model)
            right = tables.get(request.right_model)
            if left is None or right is None or request.left_model == request.right_model:
                raise AppError("RELATIONSHIP_MODEL_INVALID", "关系必须连接两个现有模型。", 422)
            for table, column in (
                (left, request.left_column),
                (right, request.right_column),
            ):
                if column not in {item["name"] for item in table["columns"]}:
                    raise AppError(
                        "RELATIONSHIP_COLUMN_INVALID", "关系字段不存在。", 422
                    )
            publication = self._publication(workspace_id, base_publication)
            self._validate_relationship_data(
                Path(publication["database_path"]), left, right, request
            )
            digest = hashlib.sha256(
                "\0".join(
                    (
                        request.left_model,
                        request.left_column,
                        request.right_model,
                        request.right_column,
                    )
                ).encode()
            ).hexdigest()[:12]
            definition = RelationshipDefinition(
                name=f"rel_{digest}",
                models=(request.left_model, request.right_model),
                join_type=request.join_type,
                condition=(
                    f"{request.left_model}.{request.left_column} = "
                    f"{request.right_model}.{request.right_column}"
                ),
            )
            relationships = {
                item.name: item
                for item in self._relationships(workspace_id, base_publication)
            }
            relationships[definition.name] = definition
            return self._prepare_candidate(
                workspace_id,
                base_publication,
                datasets,
                tuple(sorted(relationships.values(), key=lambda item: item.name)),
            )

    def _prepare_candidate(
        self,
        workspace_id: str,
        base_publication: str | None,
        datasets: tuple[DatasetRecord, ...],
        relationships: tuple[RelationshipDefinition, ...],
    ) -> PublicationCandidate:
        signature = hashlib.sha256(
            json.dumps(
                [
                    _PUBLISHER_VERSION,
                    [(item.dataset_id, item.content_hash) for item in datasets],
                    [item.model_dump(mode="json") for item in relationships],
                ],
                ensure_ascii=False,
                separators=(",", ":"),
            ).encode()
        ).hexdigest()
        active = self._publication(workspace_id, base_publication) if base_publication else None
        if active is not None and active["signature"] == signature:
            project_path = Path(active["project_path"])
            return PublicationCandidate(
                workspace_id=workspace_id,
                publication_id=base_publication,
                base_publication_id=base_publication,
                database_path=Path(active["database_path"]),
                project_path=project_path,
                project_relative_path=project_path.relative_to(self.published_root),
                signature=signature,
                datasets=datasets,
                relationships=relationships,
                changed=False,
            )

        publication_id = "p_" + signature[:24]
        project_relative = Path(workspace_id) / publication_id
        project_path = self.published_root / project_relative
        database_path = project_path / "data.duckdb"
        parsed = self._build_publication(
            workspace_id, publication_id, datasets, project_path, relationships
        )
        return PublicationCandidate(
            workspace_id=workspace_id,
            publication_id=publication_id,
            base_publication_id=base_publication,
            database_path=database_path,
            project_path=project_path,
            project_relative_path=project_relative,
            signature=signature,
            datasets=parsed,
            relationships=relationships,
        )

    def activate(self, candidate: PublicationCandidate) -> PublicationResult:
        if not candidate.changed:
            return PublicationResult(
                workspace_id=candidate.workspace_id,
                publication_id=candidate.publication_id,
                changed=False,
                datasets=_summaries(candidate.datasets),
                relationships=candidate.relationships,
            )
        with self._lock, self._connect() as connection:
            row = connection.execute(
                "SELECT active_publication_id FROM dataset_workspaces WHERE workspace_id=?",
                (candidate.workspace_id,),
            ).fetchone()
            current = row[0] if row else None
            if current != candidate.base_publication_id:
                raise AppError(
                    "PUBLICATION_CONFLICT", "数据空间已被其他上传更新，请重新上传。", 409
                )
            now = dt.datetime.now(dt.timezone.utc).isoformat()
            connection.execute(
                "DELETE FROM datasets WHERE workspace_id=?",
                (candidate.workspace_id,),
            )
            for dataset in candidate.datasets:
                connection.execute(
                    """
                    INSERT INTO datasets(
                        workspace_id,dataset_id,source_name,content_hash,source_path,
                        schema_json,updated_at
                    ) VALUES(?,?,?,?,?,?,?)
                    ON CONFLICT(workspace_id,dataset_id) DO UPDATE SET
                        source_name=excluded.source_name,
                        content_hash=excluded.content_hash,
                        source_path=excluded.source_path,
                        schema_json=excluded.schema_json,
                        updated_at=excluded.updated_at
                    """,
                    (
                        candidate.workspace_id,
                        dataset.dataset_id,
                        dataset.source_name,
                        dataset.content_hash,
                        str(dataset.source_path),
                        json.dumps(dataset.tables, ensure_ascii=False),
                        now,
                    ),
                )
            connection.execute(
                """
                INSERT OR IGNORE INTO dataset_publications(
                    workspace_id,publication_id,database_path,project_path,signature,
                    datasets_json,created_at,relationships_json
                ) VALUES(?,?,?,?,?,?,?,?)
                """,
                (
                    candidate.workspace_id,
                    candidate.publication_id,
                    str(candidate.database_path),
                    str(candidate.project_path),
                    candidate.signature,
                    json.dumps(
                        [
                            dataset.model_dump(mode="json")
                            for dataset in candidate.datasets
                        ],
                        ensure_ascii=False,
                    ),
                    now,
                    json.dumps(
                        [item.model_dump(mode="json") for item in candidate.relationships],
                        ensure_ascii=False,
                    ),
                ),
            )
            connection.execute(
                """
                INSERT INTO dataset_workspaces(workspace_id,active_publication_id) VALUES(?,?)
                ON CONFLICT(workspace_id) DO UPDATE SET
                    active_publication_id=excluded.active_publication_id
                """,
                (candidate.workspace_id, candidate.publication_id),
            )
        return PublicationResult(
            workspace_id=candidate.workspace_id,
            publication_id=candidate.publication_id,
            changed=True,
            datasets=_summaries(candidate.datasets),
            relationships=candidate.relationships,
        )

    def publications(self, workspace_id: str) -> tuple[PublicationSummary, ...]:
        if not _ID.fullmatch(workspace_id):
            raise AppError("WORKSPACE_INVALID", "数据空间标识不合法。", 422)
        with self._connect() as connection:
            active_row = connection.execute(
                "SELECT active_publication_id FROM dataset_workspaces WHERE workspace_id=?",
                (workspace_id,),
            ).fetchone()
            rows = _all_dicts(
                connection.execute(
                    "SELECT publication_id,datasets_json,relationships_json,created_at "
                    "FROM dataset_publications WHERE workspace_id=? "
                    "ORDER BY created_at DESC, publication_id DESC",
                    (workspace_id,),
                )
            )
        active = active_row[0] if active_row else None
        return tuple(
            PublicationSummary(
                publication_id=row["publication_id"],
                created_at=row["created_at"],
                active=row["publication_id"] == active,
                datasets=_summaries(
                    tuple(
                        DatasetRecord.model_validate(item)
                        for item in json.loads(row["datasets_json"] or "[]")
                    )
                ),
                relationships=tuple(
                    RelationshipDefinition.model_validate(item)
                    for item in json.loads(row["relationships_json"] or "[]")
                ),
            )
            for row in rows
        )

    def active(self, workspace_id: str) -> PublicationResult | None:
        datasets, publication_id = self._current(workspace_id)
        if not publication_id:
            return None
        return PublicationResult(
            workspace_id=workspace_id,
            publication_id=publication_id,
            changed=False,
            datasets=_summaries(datasets),
            relationships=self._relationships(workspace_id, publication_id),
        )

    def rollback(self, workspace_id: str, publication_id: str) -> PublicationResult:
        if not _ID.fullmatch(workspace_id) or not _ID.fullmatch(publication_id):
            raise AppError("PUBLICATION_INVALID", "发布版本标识不合法。", 422)
        with self._lock, self._connect() as connection:
            publication = _one_dict(
                connection.execute(
                    "SELECT * FROM dataset_publications "
                    "WHERE workspace_id=? AND publication_id=?",
                    (workspace_id, publication_id),
                )
            )
            if publication is None:
                raise AppError("PUBLICATION_NOT_FOUND", "发布版本不存在。", 404)
            if not publication["datasets_json"]:
                raise AppError(
                    "PUBLICATION_SNAPSHOT_MISSING",
                    "该历史发布缺少数据集快照，不能自动回滚。",
                    409,
                )
            datasets = tuple(
                DatasetRecord.model_validate(item)
                for item in json.loads(publication["datasets_json"])
            )
            relationships = tuple(
                RelationshipDefinition.model_validate(item)
                for item in json.loads(publication["relationships_json"] or "[]")
            )
            missing = [
                dataset.source_name
                for dataset in datasets
                if not dataset.source_path.is_file()
            ]
            project_path = Path(publication["project_path"])
            if (
                missing
                or not Path(publication["database_path"]).is_file()
                or not (project_path / "wren_project.yml").is_file()
            ):
                raise AppError(
                    "PUBLICATION_FILES_MISSING",
                    "历史发布文件不完整，不能回滚。",
                    409,
                )
            current_row = connection.execute(
                "SELECT active_publication_id FROM dataset_workspaces WHERE workspace_id=?",
                (workspace_id,),
            ).fetchone()
            current = current_row[0] if current_row else None
            if current == publication_id:
                return PublicationResult(
                    workspace_id=workspace_id,
                    publication_id=publication_id,
                    changed=False,
                    datasets=_summaries(datasets),
                    relationships=relationships,
                )
            connection.execute("DELETE FROM datasets WHERE workspace_id=?", (workspace_id,))
            now = dt.datetime.now(dt.timezone.utc).isoformat()
            for dataset in datasets:
                connection.execute(
                    """
                    INSERT INTO datasets(
                        workspace_id,dataset_id,source_name,content_hash,source_path,
                        schema_json,updated_at
                    ) VALUES(?,?,?,?,?,?,?)
                    """,
                    (
                        workspace_id,
                        dataset.dataset_id,
                        dataset.source_name,
                        dataset.content_hash,
                        str(dataset.source_path),
                        json.dumps(dataset.tables, ensure_ascii=False),
                        now,
                    ),
                )
            connection.execute(
                "UPDATE dataset_workspaces SET active_publication_id=? WHERE workspace_id=?",
                (publication_id, workspace_id),
            )
        return PublicationResult(
            workspace_id=workspace_id,
            publication_id=publication_id,
            changed=True,
            datasets=_summaries(datasets),
            relationships=relationships,
        )

    def _current(self, workspace_id: str) -> tuple[tuple[DatasetRecord, ...], str | None]:
        with self._connect() as connection:
            rows = _all_dicts(
                connection.execute(
                    "SELECT * FROM datasets WHERE workspace_id=? ORDER BY dataset_id",
                    (workspace_id,),
                )
            )
            workspace = connection.execute(
                "SELECT active_publication_id FROM dataset_workspaces WHERE workspace_id=?",
                (workspace_id,),
            ).fetchone()
        records = tuple(
            DatasetRecord(
                dataset_id=row["dataset_id"],
                source_name=row["source_name"],
                content_hash=row["content_hash"],
                source_path=Path(row["source_path"]),
                tables=tuple(json.loads(row["schema_json"])),
            )
            for row in rows
        )
        return records, workspace[0] if workspace else None

    def _publication(self, workspace_id: str, publication_id: str):
        with self._connect() as connection:
            return _one_dict(
                connection.execute(
                    "SELECT * FROM dataset_publications "
                    "WHERE workspace_id=? AND publication_id=?",
                    (workspace_id, publication_id),
                )
            )

    def _relationships(
        self, workspace_id: str, publication_id: str | None
    ) -> tuple[RelationshipDefinition, ...]:
        if publication_id is None:
            return ()
        publication = self._publication(workspace_id, publication_id)
        if publication is None:
            return ()
        return tuple(
            RelationshipDefinition.model_validate(item)
            for item in json.loads(publication["relationships_json"] or "[]")
        )

    @staticmethod
    def _validate_relationship_data(
        database_path: Path,
        left: dict[str, Any],
        right: dict[str, Any],
        request: RelationshipInput,
    ) -> None:
        left_table = _quoted(left["table"])
        right_table = _quoted(right["table"])
        left_column = _quoted(request.left_column)
        right_column = _quoted(request.right_column)
        with duckdb.connect(str(database_path), read_only=True) as connection:
            left_count, left_distinct = connection.execute(
                f"SELECT COUNT({left_column}),COUNT(DISTINCT {left_column}) "
                f"FROM {left_table}"
            ).fetchone()
            right_count, right_distinct = connection.execute(
                f"SELECT COUNT({right_column}),COUNT(DISTINCT {right_column}) "
                f"FROM {right_table}"
            ).fetchone()
            matches = connection.execute(
                f"SELECT COUNT(*) FROM {left_table} l JOIN {right_table} r "
                f"ON l.{left_column}=r.{right_column}"
            ).fetchone()[0]
        if matches == 0:
            raise AppError(
                "RELATIONSHIP_NO_MATCH", "两个字段没有可验证的匹配数据。", 422
            )
        left_unique = left_count == left_distinct
        right_unique = right_count == right_distinct
        valid = {
            "ONE_TO_ONE": left_unique and right_unique,
            "ONE_TO_MANY": left_unique,
            "MANY_TO_ONE": right_unique,
            "MANY_TO_MANY": True,
        }[request.join_type]
        if not valid:
            raise AppError(
                "RELATIONSHIP_CARDINALITY_INVALID",
                "字段唯一性与选择的关系类型不一致。",
                422,
            )

    def _store_upload(self, workspace_id: str, upload: WorkbookUpload) -> DatasetRecord:
        filename = Path(upload.filename).name
        if filename != upload.filename or Path(filename).suffix.lower() != ".xlsx":
            raise AppError("EXCEL_FILE_INVALID", "只接受不含路径的 .xlsx 文件名。", 422)
        if len(upload.content) > self.max_upload_bytes:
            raise AppError("EXCEL_TOO_LARGE", "Excel 文件超过允许大小。", 413)
        dataset_id = upload.dataset_id or _stable_dataset_id(workspace_id, filename)
        if not _ID.fullmatch(dataset_id):
            raise AppError("DATASET_ID_INVALID", "数据集标识不合法。", 422)
        digest = hashlib.sha256(upload.content).hexdigest()
        directory = self.source_root / workspace_id / dataset_id
        directory.mkdir(parents=True, exist_ok=True)
        path = directory / f"{digest}.xlsx"
        if not path.exists():
            temporary = directory / f".upload-{uuid.uuid4().hex[:12]}.tmp"
            temporary.write_bytes(upload.content)
            os.replace(temporary, path)
        return DatasetRecord(
            dataset_id=dataset_id,
            source_name=filename,
            content_hash=digest,
            source_path=path,
            tables=(),
        )

    def _build_publication(
        self,
        workspace_id: str,
        publication_id: str,
        datasets: tuple[DatasetRecord, ...],
        project_path: Path,
        relationships: tuple[RelationshipDefinition, ...] = (),
    ) -> tuple[DatasetRecord, ...]:
        if project_path.exists():
            with self._connect() as connection:
                row = connection.execute(
                    "SELECT datasets_json FROM dataset_publications "
                    "WHERE workspace_id=? AND publication_id=?",
                    (workspace_id, publication_id),
                ).fetchone()
            if row:
                return tuple(
                    DatasetRecord.model_validate(item) for item in json.loads(row[0])
                )
            shutil.rmtree(project_path)
        project_path.parent.mkdir(parents=True, exist_ok=True)
        temporary = project_path.with_name(f".tmp-{uuid.uuid4().hex[:8]}")
        temporary.mkdir()
        database_path = temporary / "data.duckdb"
        models, parsed_datasets = [], []
        try:
            connection = duckdb.connect(str(database_path))
            try:
                connection.begin()
                for dataset in datasets:
                    dataset_models, schema = self._import_workbook(connection, dataset)
                    models.extend(dataset_models)
                    parsed_datasets.append(dataset.model_copy(update={"tables": schema}))
                connection.execute(
                    "CREATE TABLE _publication_meta(workspace_id VARCHAR,publication_id VARCHAR)"
                )
                connection.execute(
                    "INSERT INTO _publication_meta VALUES(?,?)",
                    (workspace_id, publication_id),
                )
                connection.commit()
            except Exception:
                connection.rollback()
                raise
            finally:
                connection.close()
            write_wren_project(
                temporary,
                workspace_id=workspace_id,
                publication_id=publication_id,
                models=models,
                relationships=[
                    item.model_dump(mode="json") for item in relationships
                ],
            )
            os.replace(temporary, project_path)
            return tuple(parsed_datasets)
        except AppError:
            if temporary.exists():
                shutil.rmtree(temporary)
            raise
        except Exception as exc:
            if temporary.exists():
                shutil.rmtree(temporary)
            raise AppError("EXCEL_IMPORT_FAILED", "Excel 无法生成查询发布。", 422) from exc

    def _import_workbook(self, connection, dataset: DatasetRecord):
        try:
            workbook = load_workbook(
                BytesIO(dataset.source_path.read_bytes()),
                read_only=True,
                data_only=False,
                keep_links=False,
            )
        except Exception as exc:
            raise AppError("EXCEL_INVALID", "Excel 文件无法读取。", 422) from exc
        models, schema, total_cells = [], [], 0
        try:
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
                # Some Excel/WPS exports contain a stale A1:A1 dimension even though
                # the worksheet XML has a full table. Resetting makes read-only mode
                # stream the actual cells instead of silently importing one column.
                if sheet.calculate_dimension() == "A1:A1":
                    sheet.reset_dimensions()
                iterator = sheet.iter_rows()
                first = next(iterator, None)
                if first is None or not any(cell.value not in (None, "") for cell in first):
                    continue
                headers = _headers(tuple(cell.value for cell in first))
                if len(headers) > self.max_columns_per_sheet:
                    raise AppError("EXCEL_COLUMNS_EXCEEDED", "工作表列数超过限制。", 422)
                rows, types = [], [None] * len(headers)
                for row_index, cells in enumerate(iterator, start=2):
                    if row_index - 1 > self.max_rows_per_sheet:
                        raise AppError("EXCEL_ROWS_EXCEEDED", "工作表行数超过限制。", 422)
                    selected = cells[: len(headers)]
                    if any(cell.data_type == "f" for cell in selected):
                        raise AppError(
                            "EXCEL_FORMULA_UNSUPPORTED", "首版不导入公式单元格，请上传结果值。", 422
                        )
                    values = tuple(cell.value for cell in selected)
                    if not any(value is not None for value in values):
                        continue
                    values += (None,) * (len(headers) - len(values))
                    rows.append(values)
                    for index, value in enumerate(values):
                        types[index] = _merge_kind(types[index], _kind(value))
                    total_cells += len(headers)
                    if total_cells > self.max_cells_per_workbook:
                        raise AppError("EXCEL_CELLS_EXCEEDED", "Excel 单元格数量超过限制。", 422)
                wren_types = tuple(value or "VARCHAR" for value in types)
                suffix = dataset.dataset_id.replace("-", "_").replace(".", "_")
                model_name = f"ds_{suffix}_sheet_{sheet_index}"
                table_name = f"t_{suffix}_sheet_{sheet_index}"
                definitions = ",".join(
                    f"{_quoted(name)} {value}"
                    for name, value in zip(headers, wren_types, strict=True)
                )
                connection.execute(f"CREATE TABLE {_quoted(table_name)}({definitions})")
                if rows:
                    marks = ",".join("?" for _ in headers)
                    connection.executemany(
                        f"INSERT INTO {_quoted(table_name)} VALUES({marks})",
                        [tuple(_duckdb_value(value) for value in row) for row in rows],
                    )
                columns = [
                    {"name": name, "type": value}
                    for name, value in zip(headers, wren_types, strict=True)
                ]
                models.append(
                    {
                        "name": model_name,
                        "description": f"{dataset.source_name} / {sheet.title}",
                        "tableReference": {"schema": "main", "table": table_name},
                        "columns": columns,
                    }
                )
                schema.append(
                    {
                        "sheet": sheet.title,
                        "model": model_name,
                        "table": table_name,
                        "row_count": len(rows),
                        "columns": columns,
                    }
                )
        finally:
            workbook.close()
        if not models:
            raise AppError("EXCEL_EMPTY", "Excel 没有可发布的工作表。", 422)
        return models, tuple(schema)
